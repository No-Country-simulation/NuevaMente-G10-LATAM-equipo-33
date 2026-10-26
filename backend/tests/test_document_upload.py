"""HTTP multipart y extracción reales; pipeline sustituido, sin SDK ni credenciales."""
import ast
from io import BytesIO
from pathlib import Path
import runpy
import socket
import sys
from types import ModuleType, SimpleNamespace
from unittest.mock import Mock

import pytest
from fastapi import FastAPI, UploadFile, HTTPException
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from backend.services.extraction_service import extraer_texto

ROOT = Path(__file__).resolve().parents[2]
TEXT = "Una VCN es una red privada virtual."
PARAMS = dict(perfil_destinatario="Principiante", formato_salida="Quiz",
              nicho_sector="General", nivel_detalle="Intermedio")


def documento(extension, vacio=False):
    salida = BytesIO()
    if extension == ".docx":
        from docx import Document
        doc = Document()
        if not vacio:
            doc.add_paragraph(TEXT)
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "Tabla VCN"
        doc.save(salida)
    elif extension == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        if not vacio:
            grupo = slide.shapes.add_group_shape()
            grupo.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = TEXT
            slide.shapes.add_table(1, 1, 0, Inches(2), Inches(4), Inches(1)).table.cell(0, 0).text = "Tabla VCN"
        doc.save(salida)
    else:
        doc = PdfWriter()
        pagina = doc.add_blank_page(width=612, height=792)
        if not vacio:
            fuente = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                       NameObject("/Subtype"): NameObject("/Type1"),
                                       NameObject("/BaseFont"): NameObject("/Helvetica")})
            pagina[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
                DictionaryObject({NameObject("/F1"): doc._add_object(fuente)})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 40 700 Td ({TEXT}) Tj ET".encode("ascii"))
            pagina[NameObject("/Contents")] = doc._add_object(stream)
        doc.write(salida)
    return salida.getvalue()


@pytest.fixture
def api(monkeypatch):
    # Permite únicamente el socket local interno del event loop de Windows.
    conectar = socket.socket.connect
    def solo_local(sock, address):
        if isinstance(address, tuple) and address[0] in {"127.0.0.1", "::1"}:
            return conectar(sock, address)
        raise AssertionError("Red externa prohibida")
    monkeypatch.setattr(socket.socket, "connect", solo_local)
    monkeypatch.setenv("GEMINI_API_KEYS", "")
    pipeline = ModuleType("backend.services.adaptation_service")
    pipeline.adaptar_contenido = Mock(return_value={"status": "exito"})
    orquestacion = ModuleType("orquestacion_agentes.main")
    orquestacion.EntradaInvalida = type("EntradaInvalida", (Exception,), {})
    orquestacion.ErrorGeneracion = type("ErrorGeneracion", (Exception,), {})
    monkeypatch.setitem(sys.modules, pipeline.__name__, pipeline)
    monkeypatch.setitem(sys.modules, orquestacion.__name__, orquestacion)
    endpoint = runpy.run_path(str(ROOT / "backend/api/adaptation.py"))
    extractor = Mock(wraps=extraer_texto)
    endpoint["adapt"].__globals__["extraer_texto"] = extractor
    app = FastAPI()
    app.include_router(endpoint["router"], prefix="/api/v1")
    with TestClient(app) as client:
        yield client, pipeline.adaptar_contenido, extractor, endpoint["adapt"]


@pytest.mark.parametrize("extension", [".pdf", ".docx", ".pptx"])
def test_documento_valido_multipart(api, extension):
    client, pipeline, extractor, _ = api
    datos = documento(extension)
    respuesta = client.post("/api/v1/adapt", data=PARAMS,
                            files={"archivo": ("material" + extension, datos)})
    assert respuesta.status_code == 200
    extractor.assert_called_once_with("material" + extension, datos)
    request = pipeline.call_args.args[0]
    assert pipeline.call_args.kwargs["documento_bytes"] == datos
    assert TEXT in request.documento_contenido
    if extension != ".pdf":
        assert "Tabla VCN" in request.documento_contenido
    assert request.documento_titulo == "material" + extension
    for campo, valor in PARAMS.items():
        assert getattr(request, campo) == valor


@pytest.mark.parametrize("extension", [".pdf", ".docx", ".pptx"])
@pytest.mark.parametrize("caso", ["vacio", "corrupto", "sin_texto"])
def test_rechazo_antes_del_pipeline(api, extension, caso):
    client, pipeline, _, _ = api
    datos = b"" if caso == "vacio" else b"archivo corrupto" if caso == "corrupto" else documento(extension, vacio=True)
    respuesta = client.post("/api/v1/adapt", data=PARAMS,
                            files={"archivo": ("material" + extension, datos)})
    assert respuesta.status_code == (413 if caso == "vacio" else 422)
    assert isinstance(respuesta.json()["detail"], str)
    pipeline.assert_not_called()


def test_limite_exacto_y_exceso(api):
    client, pipeline, _, _ = api
    from backend.core.config import settings
    limite = settings.max_upload_mb * 1024 * 1024
    assert settings.max_upload_mb == 10
    for longitud, estado in ((limite, 200), (limite + 1, 413)):
        pipeline.reset_mock()
        respuesta = client.post("/api/v1/adapt", data=PARAMS,
                                files={"archivo": ("material.txt", b"a" * longitud)})
        assert respuesta.status_code == estado
        assert pipeline.called == (estado == 200)


def test_contrato_incompleto_y_extension(api):
    client, pipeline, _, _ = api
    assert client.post("/api/v1/adapt", json=PARAMS).status_code == 422
    assert client.post("/api/v1/adapt", data={}, files={"archivo": ("a.txt", b"texto")}).status_code == 422
    assert client.post("/api/v1/adapt", data=PARAMS, files={"archivo": ("a.exe", b"texto")}).status_code == 415
    pipeline.assert_not_called()


def test_archivo_cerrado_ante_error(api):
    _, pipeline, _, adapt = api
    archivo = UploadFile(filename="mal.pdf", file=BytesIO(b"corrupto"))
    with pytest.raises(HTTPException):
        adapt(archivo=archivo, **PARAMS)
    assert archivo.file.closed
    pipeline.assert_not_called()


@pytest.mark.parametrize("backend_falla", [False, True, "http_error"])
def test_frontend_transporta_bytes_y_conserva_simulacion(api, monkeypatch, backend_falla):
    client, _, _, _ = api
    datos = documento(".docx")
    estado = {"archivo_temporal": SimpleNamespace(name="material.docx", getvalue=lambda: datos)}
    st = SimpleNamespace(session_state=estado, rerun=lambda: None)
    carga = ast.parse((ROOT / "interfaz/pages/seccion1_carga.py").read_text(encoding="utf-8"))
    ns = {"st": st}
    exec(compile(ast.Module(body=[n for n in carga.body if isinstance(n, ast.FunctionDef) and n.name == "procesar_y_continuar"], type_ignores=[]), "carga", "exec"), ns)
    ns["procesar_y_continuar"]()
    assert estado["archivo_bytes"] == datos
    assert estado["documento_contenido"] == ""
    params = ast.parse((ROOT / "interfaz/pages/seccion2_parametros.py").read_text(encoding="utf-8"))
    ns.update(FORMATOS_SELECCION=["Flashcards de Memorización"], ADAPT_URL="/api/v1/adapt", API_TIMEOUT_SECONDS=300)
    exec(compile(ast.Module(body=[n for n in params.body if isinstance(n, ast.FunctionDef) and n.name in {"generar", "_crear_resultado_simulado"}], type_ignores=[]), "parametros", "exec"), ns)
    requests = ModuleType("requests")
    def post(url, **kwargs):
        assert kwargs.pop("timeout") == 300
        if backend_falla == "http_error":
            return SimpleNamespace(status_code=422)
        return client.post(url, **kwargs)
    requests.post = Mock(side_effect=RuntimeError("offline") if backend_falla is True else post)
    monkeypatch.setitem(sys.modules, "requests", requests)
    ns["generar"]()
    assert requests.post.call_args.kwargs["files"]["archivo"][1] == datos
    assert "json" not in requests.post.call_args.kwargs
    assert estado["resultado"]["status"] == "exito" and estado["paso_actual"] == 3
    if backend_falla:
        assert estado["resultado"]["almacenamiento_oci"]["status_upload"] == "completado"
